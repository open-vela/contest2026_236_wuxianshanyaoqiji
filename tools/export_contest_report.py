"""Export the official-template technical report, including tables and diagrams."""
from pathlib import Path
import json
import re
import shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'docs/contest/技术报告.md'
DEST = ROOT/'artifacts/contest/无限闪耀绮迹-绮迹-技术报告.docx'

def font(run, size=None, bold=None):
    run.font.name='Calibri'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    if size is not None:run.font.size=Pt(size)
    if bold is not None:run.bold=bold

def main():
    raw=SOURCE.read_text(encoding='utf-8')
    abstract=raw.split('## 2、摘要\n',1)[1].split('\n## 3、正文',1)[0].strip()
    assert len(re.sub(r'\s','',abstract))<=300
    assert '无限闪耀绮迹' in raw and '张熙哲' in raw and '彭馨怡' in raw
    assert not any(name in raw for name in ('无险山妖气迹','无限山妖气姬'))
    doc=Document();s=doc.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=Cm(1.8);s.left_margin=s.right_margin=Cm(1.9)
    s.header_distance=s.footer_distance=Cm(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name]
        style.font.name='Calibri';style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        style.font.color.rgb=RGBColor.from_string('24344B')
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',24),('Heading 1',17),('Heading 2',14),('Heading 3',11.5)]:
        doc.styles[name].font.size=Pt(size)
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.core_properties.title='绮迹 · 技术报告'
    doc.core_properties.author='无限闪耀绮迹'
    doc.core_properties.subject='2026 首届 openvela AI 硬件开发者大赛'
    header=s.header.paragraphs[0];header.text='无限闪耀绮迹  /  236                                  绮迹 · 技术报告'
    for r in header.runs:font(r,8)
    footer=s.footer.paragraphs[0];footer.alignment=2
    footer.add_run('无限闪耀绮迹  ·  ')
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
    for r in footer.runs:font(r,8)
    lines=raw.splitlines();i=0;table_count=0;image_count=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip());i+=1
            rows=[[c.strip() for c in row.strip('|').split('|')] for row in block]
            rows=[row for row in rows if not all(re.fullmatch(r':?-+:?',c) for c in row)]
            table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid';table.autofit=False
            widths=([4,13.2] if len(rows[0])==2 else [3.1,5.1,9] if len(rows[0])==3 else [3.1,1.2,6.9,6])
            for rowno,values in enumerate(rows):
                cells=table.add_row().cells
                trpr=cells[0]._tc.getparent().get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if rowno==0:trpr.append(OxmlElement('w:tblHeader'))
                for index,(cell,value) in enumerate(zip(cells,values)):
                    cell.width=Cm(widths[index]);cell.vertical_alignment=1
                    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                    p.paragraph_format.line_spacing=1.08
                    if rowno==0:p.paragraph_format.keep_with_next=True
                    run=p.add_run(value);font(run,9.5,rowno==0)
                    if rowno==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF0F7');cell._tc.get_or_add_tcPr().append(sh)
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
            table_count+=1;continue
        match=re.fullmatch(r'!\[([^]]+)\]\(([^)]+)\)',line)
        if match:
            # Report image paths use repository-root artifacts for stable generation.
            path=ROOT/'artifacts/contest/template-review-assets'/Path(match[2]).name
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(path),width=Cm(17.1))
            caption=doc.add_paragraph(f'图 {image_count+1}  {match[1]}');caption.alignment=1
            for r in caption.runs:font(r,9)
            image_count+=1;continue
        m=re.match(r'^(#{1,4}) (.*)$',line)
        if m:
            depth=len(m[1]);value=m[2]
            p=doc.add_paragraph(value,'Title' if depth==1 else f'Heading {depth-1}')
            if depth==3 and value.startswith('3.2 '):
                p.paragraph_format.page_break_before=True
        else:doc.add_paragraph(line)
    DEST.parent.mkdir(parents=True,exist_ok=True);doc.save(DEST)
    check=Document(DEST)
    text='\n'.join(p.text for p in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for row in t.rows for c in row.cells)
    for section in ['1、信息表','2、摘要']+[f'3.{n} ' for n in range(1,8)]:assert section in text,section
    assert all(value in text for value in ['无限闪耀绮迹','张熙哲','彭馨怡','/data/agent/skills/','ESP-IDF'])
    assert len(check.inline_shapes)==image_count==3
    assert len(check.tables)==table_count and table_count>=14
    shutil.copy2(DEST,ROOT/'artifacts/contest/绮迹-参赛作品介绍.docx')
    report={'template_sections':'PASS','abstract_characters':len(re.sub(r'\s','',abstract)),
            'tables':len(check.tables),'diagrams':len(check.inline_shapes),'team':'无限闪耀绮迹',
            'members':['张熙哲','彭馨怡'],'photos_included':False,'runtime_skill':'Deployment/trigger evidence not provided',
            'unknown_metrics':'Explicitly marked unmeasured/unavailable','document':DEST.name}
    (ROOT/'artifacts/contest/template-report-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,ensure_ascii=True))

if __name__=='__main__':main()
