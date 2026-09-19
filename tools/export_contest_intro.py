"""Export the maintained contest introduction to a standalone Word document."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent.parent
source=root/'docs/contest/作品介绍.md'
destination=root/'artifacts/contest/绮迹-参赛作品介绍.docx'
doc=Document()
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=section.bottom_margin=Cm(2)
section.left_margin=section.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2']:
    style=doc.styles[name]
    style.font.name='Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.25
doc.styles['Title'].font.color.rgb=RGBColor.from_string('55468B')
doc.styles['Heading 1'].font.size=Pt(14)
doc.styles['Heading 1'].font.color.rgb=RGBColor.from_string('55468B')
for line in source.read_text(encoding='utf-8').splitlines():
    if not line.strip():continue
    line=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1（\2）',line)
    line=line.replace('`','')
    if line.startswith('# '):doc.add_heading(line[2:],0)
    elif line.startswith('## '):doc.add_heading(line[3:],1)
    else:doc.add_paragraph(line)
footer=section.footer.paragraphs[0]
footer.alignment=2
footer.add_run('绮迹 · 参赛介绍草稿 | ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
destination.parent.mkdir(parents=True,exist_ok=True)
doc.save(destination)
check=Document(destination)
assert '待验收' in '\n'.join(p.text for p in check.paragraphs)
assert len(check.paragraphs)>15
print(destination)
